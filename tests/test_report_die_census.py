"""Analyzable-die census: unique Site+Slot, not loops / bins / board rows."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from lpddr_fail_analyzer.classify import BAD_COLUMN, analyze_all_dies, classify_unique_cells
from lpddr_fail_analyzer.ingest import load_board_census, load_fails
from lpddr_fail_analyzer.pipeline import run_analyze
from lpddr_fail_analyzer.report import (
    BIN_COUNT_TIP,
    BOARD_MISSING_NOTE,
    MULTI_LOOP_DIE_TIP,
    NO_STRUCTURE_MARK,
    SECTION_ANALYZABLE,
    SECTION_BOARD_ONLY,
)

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_XLSX = ROOT / "samples" / "sample_batch.xlsx"

CSV_HEADER = (
    "Site,Slot,Loop,Pattern Name,Linear ADDR,ROW,BANK,COL,EXP Value,RD Value,"
    "Re-read value1,Re-read value2,Re-read value3,XOR Val1\n"
)


def _write_csv(path: Path, body: str) -> Path:
    path.write_text(CSV_HEADER + body, encoding="utf-8")
    return path


def _col_line_cells(n: int = 8) -> list[tuple[int, int, int]]:
    return [(0x100 + r, 0, 0x10) for r in range(n)]


def _fail_body_rows(site: int, slot: int, loop: int, cells: list[tuple[int, int, int]]) -> str:
    lines = []
    for i, (row, bank, col) in enumerate(cells):
        ident = f"{site},{slot},{loop},pA" if i == 0 else ",,,"
        lines.append(
            f"{ident},0x{i:X},0x{row:X},{bank},0x{col:X},00,FF,FF,FF,FF,80\n"
        )
    return "".join(lines)


def test_multi_loop_same_site_slot_counts_as_one_analyzable_die(tmp_path: Path):
    cells = _col_line_cells(8)
    body = "".join(_fail_body_rows(15, 7, loop, cells) for loop in (1, 2, 3))
    csv_path = _write_csv(tmp_path / "three_loops.csv", body)

    df, _, stats = load_fails(csv_path)
    assert stats.dropped_rows == 0
    assert len(df) == 24
    dies = analyze_all_dies(df)
    assert len(dies) == 1
    die = dies[0]
    assert die.site == "15"
    assert die.slot == "7"
    assert die.n_loops == 3
    assert die.n_fail_rows == 24
    assert die.n_unique_cells == 8
    assert die.primary_label == BAD_COLUMN

    once = classify_unique_cells(cells)
    triple = classify_unique_cells(cells * 3)
    assert once["flags"] == triple["flags"]
    assert once["evidence"]["n_unique_cells"] == 8
    assert triple["evidence"]["n_unique_cells"] == 8
    assert once["evidence"]["top_col_share"] == triple["evidence"]["top_col_share"]

    out = tmp_path / "three-loops-out"
    result = run_analyze(csv_path, out)
    assert result.n_analyzable_dies == 1
    text = (out / "report.md").read_text(encoding="utf-8")
    assert f"- 可分析颗数（fail_msg 唯一 Site+Slot）： **1**" in text
    assert f"## {SECTION_ANALYZABLE}" in text
    assert f"## {SECTION_BOARD_ONLY}" in text
    assert MULTI_LOOP_DIE_TIP in text
    assert BIN_COUNT_TIP in text
    assert BOARD_MISSING_NOTE in text
    summary = pd.read_csv(out / "summary.csv")
    analyzable = summary[summary["section"] == SECTION_ANALYZABLE]
    assert len(analyzable) == 1
    assert int(analyzable.iloc[0]["n_loops"]) == 3
    assert int(analyzable.iloc[0]["n_unique_cells"]) == 8
    assert int(analyzable.iloc[0]["n_analyzable_dies"]) == 1


def test_two_site_slots_are_two_dies_even_with_loops(tmp_path: Path):
    cells = _col_line_cells(8)
    body = _fail_body_rows(15, 7, 1, cells) + _fail_body_rows(15, 7, 2, cells)
    body += _fail_body_rows(1, 1, 1, cells)
    csv_path = _write_csv(tmp_path / "two_dies.csv", body)
    df, _, _ = load_fails(csv_path)
    dies = analyze_all_dies(df)
    keys = sorted((d.site, d.slot) for d in dies)
    assert keys == [("1", "1"), ("15", "7")]
    by_key = {(d.site, d.slot): d for d in dies}
    assert by_key[("15", "7")].n_loops == 2
    assert by_key[("1", "1")].n_loops == 1
    assert all(d.n_unique_cells == 8 for d in dies)


def _xlsx_board_and_fail(path: Path) -> Path:
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum["A3"] = "PN"
    ws_sum["B3"] = "DEMO-LPDDR4-PN"
    ws_sum["A4"] = "LOT ID"
    ws_sum["B4"] = "LOT.BOARD"

    board = wb.create_sheet("board_msg")
    board["A1"] = "LOT ID *"
    board["B1"] = "PN"
    board["C1"] = "Board ID"
    board["E1"] = "Result"
    board["F1"] = "Bin Code"
    board["C2"] = "Site"
    board["D2"] = "Slot"
    # 15/7 fail three board rounds (still ONE die); 1/1 fail no dump; 2/2 pass only
    rows = [
        ("LOT.BOARD", "DEMO", 15, 7, "fail", "05"),
        ("LOT.BOARD", "DEMO", 15, 7, "fail", "05"),
        ("LOT.BOARD", "DEMO", 15, 7, "fail", "05"),
        ("LOT.BOARD", "DEMO", 1, 1, "fail", "03"),
        ("LOT.BOARD", "DEMO", 1, 1, "fail", "03"),
        ("LOT.BOARD", "DEMO", 2, 2, "pass", "01"),
    ]
    for i, rec in enumerate(rows, start=3):
        for col, val in enumerate(rec, start=1):
            board.cell(i, col, val)

    fail = wb.create_sheet("fail_msg")
    fail["A1"] = "Board ID"
    fail["C1"] = "Fail Information"
    headers = [
        "Site",
        "Slot",
        "Loop",
        "Pattern Name",
        "Linear ADDR",
        "ROW",
        "BANK",
        "COL",
        "EXP Value",
        "RD Value",
        "Re-read value1",
        "Re-read value2",
        "Re-read value3",
        "XOR Val1",
    ]
    for col, name in enumerate(headers, 1):
        fail.cell(3, col, name)
    cells = _col_line_cells(8)
    r = 4
    for loop in (1, 2, 3):
        for i, (row, bank, col) in enumerate(cells):
            if i == 0:
                fail.cell(r, 1, 15)
                fail.cell(r, 2, 7)
                fail.cell(r, 3, loop)
                fail.cell(r, 4, "pA")
            fail.cell(r, 5, f"0x{i:X}")
            fail.cell(r, 6, f"0x{row:X}")
            fail.cell(r, 7, bank)
            fail.cell(r, 8, f"0x{col:X}")
            fail.cell(r, 9, "00")
            fail.cell(r, 10, "FF")
            fail.cell(r, 11, "FF")
            fail.cell(r, 12, "FF")
            fail.cell(r, 13, "FF")
            fail.cell(r, 14, "80")
            r += 1
    wb.save(path)
    return path


def test_board_no_dump_section_from_xlsx(tmp_path: Path):
    xlsx = _xlsx_board_and_fail(tmp_path / "board_mix.xlsx")
    board = load_board_census(xlsx)
    assert board.present
    assert board.n_fail_rows == 5
    assert board.n_unique_fail_dies == 2
    keys = {(d.site, d.slot) for d in board.fail_dies}
    assert keys == {("15", "7"), ("1", "1")}

    result = run_analyze(xlsx, tmp_path / "board-mix-out")
    assert result.n_analyzable_dies == 1
    assert result.n_board_no_dump == 1
    text = (result.out_dir / "report.md").read_text(encoding="utf-8")
    assert f"## {SECTION_ANALYZABLE}" in text
    assert f"## {SECTION_BOARD_ONLY}" in text
    assert MULTI_LOOP_DIE_TIP in text
    assert BIN_COUNT_TIP in text
    assert NO_STRUCTURE_MARK in text
    assert f"- 可分析颗数（fail_msg 唯一 Site+Slot）： **1**" in text
    assert "- 仅 board 不良、无 dump： **1**" in text

    start = text.index(f"## {SECTION_BOARD_ONLY}")
    rest = text[start:]
    end = rest.find("\n## ", 4)
    board_section = rest if end < 0 else rest[:end]
    assert "| 1 | 1 |" in board_section
    assert "| 15 | 7 |" not in board_section
    assert "| 2 | 2 |" not in board_section
    assert NO_STRUCTURE_MARK in board_section

    summary = pd.read_csv(result.out_dir / "summary.csv")
    assert set(summary["section"]) == {SECTION_ANALYZABLE, SECTION_BOARD_ONLY}
    only = summary[summary["section"] == SECTION_BOARD_ONLY]
    assert len(only) == 1
    assert str(int(only.iloc[0]["site"])) == "1"
    assert str(int(only.iloc[0]["slot"])) == "1"
    assert only.iloc[0]["note"] == NO_STRUCTURE_MARK


def test_xlsx_without_board_sheet_says_so(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "fail_msg"
    headers = [
        "Site",
        "Slot",
        "Loop",
        "Pattern Name",
        "Linear ADDR",
        "ROW",
        "BANK",
        "COL",
    ]
    for col, name in enumerate(headers, 1):
        ws.cell(1, col, name)
    ws.cell(2, 1, 9)
    ws.cell(2, 2, 8)
    ws.cell(2, 3, 1)
    ws.cell(2, 4, "p")
    ws.cell(2, 5, "0x1")
    ws.cell(2, 6, "0x10")
    ws.cell(2, 7, 0)
    ws.cell(2, 8, "0x0")
    path = tmp_path / "no_board.xlsx"
    wb.save(path)
    board = load_board_census(path)
    assert not board.present
    assert board.missing_reason == "no_sheet"
    text = (run_analyze(path, tmp_path / "no-board-out").out_dir / "report.md").read_text(
        encoding="utf-8"
    )
    assert f"## {SECTION_BOARD_ONLY}" in text
    assert BOARD_MISSING_NOTE in text


def test_sample_batch_splits_analyzable_and_board_only(tmp_path: Path):
    result = run_analyze(SAMPLE_XLSX, tmp_path / "sample-census")
    assert result.run_status == "ok"
    assert result.n_analyzable_dies == 1
    assert result.n_board_no_dump == 19
    text = (result.out_dir / "report.md").read_text(encoding="utf-8")
    assert f"## {SECTION_ANALYZABLE}" in text
    assert f"## {SECTION_BOARD_ONLY}" in text
    assert MULTI_LOOP_DIE_TIP in text
    assert BIN_COUNT_TIP in text
    assert NO_STRUCTURE_MARK in text
    assert f"- 可分析颗数（fail_msg 唯一 Site+Slot）： **1**" in text
    assert "- 仅 board 不良、无 dump： **19**" in text
    assert "有地址明细、可分析颗粒 1 颗" in text

    start = text.index(f"## {SECTION_BOARD_ONLY}")
    rest = text[start:]
    end = rest.find("\n## ", 4)
    board_section = rest if end < 0 else rest[:end]
    assert "| 15 | 7 |" not in board_section
    assert "| 1 | 7 |" in board_section
    assert board_section.count(NO_STRUCTURE_MARK) >= 19

    summary = pd.read_csv(result.out_dir / "summary.csv")
    analyzable = summary[summary["section"] == SECTION_ANALYZABLE]
    board_only = summary[summary["section"] == SECTION_BOARD_ONLY]
    assert len(analyzable) == 1
    assert str(int(analyzable.iloc[0]["site"])) == "15"
    assert str(int(analyzable.iloc[0]["slot"])) == "7"
    assert int(analyzable.iloc[0]["n_loops"]) == 3
    assert len(board_only) == 19
    assert "bad_column" in set(analyzable["primary_label"])
    assert all(board_only["note"] == NO_STRUCTURE_MARK)
