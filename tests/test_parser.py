from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from lpddr_fail_analyzer.ingest import find_header_row, load_fails, normalize_fail_frame


def _tiny_xlsx(path: Path) -> Path:
    wb = Workbook()
    # Title-only first sheet name matching production layout.
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum["A1"] = "Handler 结批汇总"
    ws_sum["A3"] = "PN"
    ws_sum["B3"] = "DEMO-LPDDR4X-PN"
    ws_sum["A4"] = "LOT ID"
    ws_sum["B4"] = "LOT.FFWD"
    ws_sum["A5"] = "测试温度"
    ws_sum["B5"] = "85"
    ws_sum["E7"] = "VDD1"
    ws_sum["F7"] = "1800mV"
    ws_sum["E8"] = "VDD2"
    ws_sum["F8"] = "1100mV"
    ws_sum["E9"] = "VDDQ"
    ws_sum["F9"] = "1100mV"

    ws = wb.create_sheet("fail_msg")
    ws["A1"] = "Board ID"
    ws["C1"] = "Fail Information"
    ws.append([])  # row 2 leftover title space
    headers = [
        "Site",
        "Slot",
        "Loop",
        "Pattern Name",
        "Linear ADDR",
        "ROW",
        "BANK",
        "COL",
        "EXP Value",
        "RD Value",
        "Re-read value1",
        "Re-read value2",
        "Re-read value3",
        "XOR Val1",
    ]
    for col, name in enumerate(headers, 1):
        ws.cell(3, col, name)

    def write_row(r, values):
        for c, v in enumerate(values, 1):
            if v is not None:
                ws.cell(r, c, v)

    # Die A first fail
    write_row(
        4,
        [1, 2, 1, "movinv8", "0x10", "0x100", 0, "0x1", "00", "FF", "FF", "FF", "FF", "FF"],
    )
    # Same die, identity blank (must fill-forward)
    write_row(
        5,
        [None, None, None, None, "0x20", "0x200", 0, "0x1", "00", "FF", "FF", "FF", "FF", "FF"],
    )
    # Pattern changes; Site/Slot still blank
    write_row(
        6,
        [None, None, 1, "btf", "0x30", "0x300", 0, "0x1", "00", "FF", "FF", "FF", "FF", "FF"],
    )
    # Die B
    write_row(
        7,
        [3, 4, 2, "isi", "0x40", "0x111", 1, "0x22", "00", "01", "01", "01", "01", "01"],
    )
    write_row(
        8,
        [None, None, None, None, "0x50", "0x112", 1, "0x22", "00", "01", "01", "01", "01", "01"],
    )
    wb.save(path)
    return path


def test_find_header_row_skips_titles():
    raw = pd.DataFrame(
        [
            ["Board ID", None, "Fail Information"],
            [
                "Site",
                "Slot",
                "Loop",
                "Pattern Name",
                "Linear ADDR",
                "ROW",
                "BANK",
                "COL",
            ],
            [1, 2, 1, "p", "0x1", "0x1", 0, "0x0"],
        ]
    )
    idx, mapping = find_header_row(raw.itertuples(index=False, name=None))
    assert idx == 1
    assert mapping[0] == "site"
    assert mapping[7] == "col"


def test_csv_fill_forward(tmp_path: Path):
    csv_path = tmp_path / "blank.csv"
    csv_path.write_text(
        "Site,Slot,Loop,Pattern Name,Linear ADDR,ROW,BANK,COL,EXP Value,RD Value,"
        "Re-read value1,Re-read value2,Re-read value3,XOR Val1\n"
        "9,8,1,foo,0x1,0x10,0,0x0,0,1,1,1,1,1\n"
        ",,, ,0x2,0x11,0,0x0,0,1,1,1,1,1\n",
        encoding="utf-8",
    )
    df, meta = load_fails(csv_path)
    assert list(df["site"].unique()) == ["9"]
    assert list(df["slot"].unique()) == ["8"]
    assert list(df["pattern_name"].unique()) == ["foo"]
    assert len(df) == 2
    assert meta.lot_id is None


def test_xlsx_fill_forward_and_summary(tmp_path: Path):
    xlsx = _tiny_xlsx(tmp_path / "tiny.xlsx")
    df, meta = load_fails(xlsx)
    assert meta.pn == "DEMO-LPDDR4X-PN"
    assert meta.lot_id == "LOT.FFWD"
    assert meta.temperature == "85"
    assert meta.vdd1 == "1800mV"
    assert meta.vdd2 == "1100mV"
    assert meta.vddq == "1100mV"

    groups = sorted(df.groupby(["site", "slot"]).groups)
    assert groups == [("1", "2"), ("3", "4")]
    die_a = df[(df["site"] == "1") & (df["slot"] == "2")]
    assert len(die_a) == 3
    assert set(die_a["pattern_name"]) == {"movinv8", "btf"}
    assert list(die_a["loop"].unique()) == ["1"]
    die_b = df[(df["site"] == "3") & (df["slot"] == "4")]
    assert len(die_b) == 2
    assert list(die_b["pattern_name"].unique()) == ["isi"]


def test_normalize_accepts_already_headed_frame():
    raw = pd.DataFrame(
        [
            [
                "Site",
                "Slot",
                "Loop",
                "Pattern Name",
                "Linear ADDR",
                "ROW",
                "BANK",
                "COL",
            ],
            [5, 6, 1, "p", "0xAA", "0x1", 7, "0x2"],
        ]
    )
    out = normalize_fail_frame(raw)
    assert out.iloc[0]["site"] == "5"
    assert out.iloc[0]["row_i"] == 0x1
    assert out.iloc[0]["col_i"] == 0x2
    assert out.iloc[0]["bank_i"] == 7
